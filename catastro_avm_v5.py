"""
AVM - Consulta de datos catastrales + precio de mercado Idealista
Uso:
    python catastro_avm_v5.py referencias.xlsx
    python catastro_avm_v5.py 2730701VK4723B0016PH
El Excel de referencias puede tener columna 'tipo' (piso/chalet/parking/local)

Mejoras v5:
  - Filtrado IQR de outliers para todos los tipos (no solo parking)
  - Mediana en lugar de media para el precio central
  - Paginacion Idealista hasta 3 paginas o >=15 comparables
  - Filtro de m2 usa m2_vivienda en lugar de superficie total
  - Autodeteccion de tipo desde uso_principal del catastro
  - Circulo geografico con 12 puntos (antes 6)
  - Coeficiente de variacion para indicar fiabilidad del dato
  - Flag de aviso para planta baja / sotano
"""

import sys
import re
import xml.etree.ElementTree as ET
import statistics
from urllib.request import urlopen, Request
from urllib.error import URLError, HTTPError
from urllib.parse import urlencode, quote
from datetime import datetime
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "https://ovc.catastro.meh.es/ovcservweb/OVCSWLocalizacionRC"
HEADERS = {"User-Agent": "Mozilla/5.0", "Accept": "text/xml"}

# Mapeo uso_principal catastro -> tipo Idealista (autodeteccion)
USO_A_TIPO = {
    "APARCAMIENTO":            "parking",
    "ALMACEN-ESTACIONAMIENTO": "parking",
    "COMERCIAL":               "local",
    "INDUSTRIAL":              "local",
    "ALMACEN":                 "local",
    "OFICINAS":                "local",
    "OCIO Y HOSTELERIA":       "local",
    "SANIDAD Y BENEFICENCIA":  "local",
    "CULTURAL":                "local",
    "RELIGIOSO":               "local",
    "ESPECTACULOS":            "local",
    "DEPORTIVO":               "local",
}

COLUMNAS = [
    ("referencia_catastral",   "Ref. Catastral"),
    ("tipo_via",               "Tipo Via"),
    ("nombre_via",             "Nombre Via"),
    ("numero",                 "Numero"),
    ("bloque",                 "Bloque"),
    ("escalera",               "Escalera"),
    ("planta",                 "Planta"),
    ("puerta",                 "Puerta"),
    ("codigo_postal",          "C.P."),
    ("municipio",              "Municipio"),
    ("provincia",              "Provincia"),
    ("comunidad",              "Comunidad Autonoma"),
    ("clase_inmueble",         "Clase Inmueble"),
    ("uso_principal",          "Uso Principal"),
    ("superficie_m2",          "Superficie (m2)"),
    ("m2_vivienda",            "m2 Vivienda"),
    ("m2_aparcamiento",        "m2 Aparcamiento"),
    ("m2_almacen",             "m2 Almacen"),
    ("m2_comercio",            "m2 Comercio"),
    ("m2_deportivo",           "m2 Deportivo"),
    ("m2_elem_comunes",        "m2 Elem. Comunes"),
    ("m2_soporte",             "m2 Soporte"),
    ("m2_oficina",             "m2 Oficina"),
    ("m2_hotelero",            "m2 Hotelero"),
    ("m2_industrial",          "m2 Industrial"),
    ("m2_sanitario",           "m2 Sanitario"),
    ("m2_cultural",            "m2 Cultural"),
    ("m2_otros",               "m2 Otros"),
    ("año_construccion",       "Ano Construccion"),
    ("num_plantas",            "No Plantas"),
    ("coef_participacion",     "Coef. Participacion (%)"),
    ("valor_catastral_total",  "Valor Catastral Total (EUR)"),
    ("valor_catastral_suelo",  "Valor Catastral Suelo (EUR)"),
    ("valor_catastral_const",  "Valor Catastral Const. (EUR)"),
    ("coord_lon_wgs84",        "Longitud"),
    ("coord_lat_wgs84",        "Latitud"),
    ("tipo_busqueda",          "Tipo Busqueda"),
    ("precio_m2_idealista",    "Precio m2 Medio (EUR)"),
    ("precio_m2_min",          "Precio m2 Min (EUR)"),
    ("precio_m2_max",          "Precio m2 Max (EUR)"),
    ("coef_variacion",         "Coef. Variacion (%)"),
    ("valoracion_idealista",   "Valoracion Idealista (EUR)"),
    ("valoracion_vivienda",    "Valoracion Vivienda (EUR)"),
    ("precio_min_comparable",  "Precio Min Comparable (EUR)"),
    ("precio_max_comparable",  "Precio Max Comparable (EUR)"),
    ("num_comparables",        "No Comparables"),
    ("radio_busqueda",         "Radio Busqueda"),
    ("planta_baja_flag",       "Aviso Planta"),
    ("error_callejero",        "Error"),
]

COLUMNAS_LINKS = [
    ("link_gmaps",     "Google Maps"),
    ("link_idealista", "Idealista"),
    ("link_catastro",  "Catastro"),
]

YELLOW_KEYS  = {"radio_busqueda"}
GREEN_KEYS   = {
    "tipo_busqueda", "precio_m2_idealista", "precio_m2_min", "precio_m2_max",
    "coef_variacion", "valoracion_idealista", "valoracion_vivienda",
    "precio_min_comparable", "precio_max_comparable", "num_comparables",
}
WARNING_KEYS = {"planta_baja_flag"}


# ---------------------------------------------------------------------------
# Utilidades catastro
# ---------------------------------------------------------------------------

def limpiar_refcat(ref):
    return re.sub(r"\s+", "", ref).upper()


def get_xml(url):
    try:
        req = Request(url, headers=HEADERS)
        with urlopen(req, timeout=10) as resp:
            return ET.fromstring(resp.read())
    except HTTPError as e:
        raise RuntimeError(f"HTTP {e.code}") from e
    except URLError as e:
        raise RuntimeError(f"Error de red: {e.reason}") from e
    except ET.ParseError as e:
        raise RuntimeError(f"XML invalido: {e}") from e


def texto(elem, *tags):
    if elem is None:
        return ""
    for tag in tags:
        for child in elem.iter():
            local = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if local.lower() == tag.lower() and child.text:
                return child.text.strip()
    return ""


def obtener_datos_callejero(refcat):
    params = urlencode({"Provincia": "", "Municipio": "", "RC": refcat})
    root = get_xml(f"{BASE}/OVCCallejero.asmx/Consulta_DNPRC?{params}")
    resultado = {k: v for k, v in {
        "tipo_via":              texto(root, "tv"),
        "nombre_via":            texto(root, "nv"),
        "numero":                texto(root, "pnp"),
        "bloque":                texto(root, "bq"),
        "escalera":              texto(root, "es"),
        "planta":                texto(root, "pt"),
        "puerta":                texto(root, "pu"),
        "codigo_postal":         texto(root, "dp"),
        "municipio":             texto(root, "nm"),
        "provincia":             texto(root, "np"),
        "comunidad":             texto(root, "nca"),
        "clase_inmueble":        texto(root, "cn"),
        "uso_principal":         texto(root, "luso"),
        "superficie_m2":         texto(root, "sfc"),
        "año_construccion":      texto(root, "ant"),
        "num_plantas":           texto(root, "npt"),
        "coef_participacion":    texto(root, "cpt"),
        "valor_catastral_total": texto(root, "vcat"),
        "valor_catastral_suelo": texto(root, "vcatsue"),
        "valor_catastral_const": texto(root, "vcatcon"),
        "delegacion":            texto(root, "cp"),
        "cod_municipio":         texto(root, "cmc"),
    }.items() if v}

    USOS_CATASTRO = {
        "VIVIENDA":          "m2_vivienda",
        "APARCAMIENTO":      "m2_aparcamiento",
        "ALMACEN":           "m2_almacen",
        "COMERCIO":          "m2_comercio",
        "DEPORTIVO":         "m2_deportivo",
        "ELEMENTOS COMUNES": "m2_elem_comunes",
        "SOPORTE":           "m2_soporte",
        "OFICINA":           "m2_oficina",
        "HOTELERO":          "m2_hotelero",
        "INDUSTRIAL":        "m2_industrial",
        "SANITARIO":         "m2_sanitario",
        "CULTURAL":          "m2_cultural",
        "RELIGIOSO":         "m2_religioso",
        "ESPECTACULOS":      "m2_espectaculos",
        "OCIO":              "m2_ocio",
        "AGRICOLA":          "m2_agricola",
    }

    m2_por_uso = {}
    todos = list(root.iter())
    for i, elem in enumerate(todos):
        tag = elem.tag.split("}")[-1] if "}" in elem.tag else elem.tag
        if tag.lower() == "lcd" and elem.text:
            lcd_val = elem.text.strip().upper()
            for j in range(i + 1, min(i + 20, len(todos))):
                stag = todos[j].tag.split("}")[-1] if "}" in todos[j].tag else todos[j].tag
                if stag.lower() == "stl" and todos[j].text:
                    try:
                        stl_val = float(todos[j].text.strip())
                        if stl_val > 0:
                            campo = None
                            for uso_key, uso_campo in USOS_CATASTRO.items():
                                if uso_key in lcd_val:
                                    campo = uso_campo
                                    break
                            if campo is None:
                                campo = "m2_otros"
                            m2_por_uso[campo] = m2_por_uso.get(campo, 0) + stl_val
                    except Exception:
                        pass
                    break

    for campo, valor in m2_por_uso.items():
        resultado[campo] = str(round(valor))

    return resultado


def obtener_coordenadas(refcat):
    params = urlencode({"Provincia": "", "Municipio": "", "SRS": "EPSG:4326", "RC": refcat[:14]})
    root = get_xml(f"{BASE}/OVCCoordenadas.asmx/Consulta_CPMRC?{params}")
    lon = texto(root, "xcen") or texto(root, "xpc")
    lat = texto(root, "ycen") or texto(root, "ypc")
    return {k: v for k, v in {"coord_lon_wgs84": lon, "coord_lat_wgs84": lat}.items() if v}


def generar_links(datos):
    lat  = datos.get("coord_lat_wgs84", "")
    lon  = datos.get("coord_lon_wgs84", "")
    cp   = datos.get("codigo_postal", "")
    mun  = datos.get("municipio", "").lower().replace(" ", "-")
    prov = datos.get("provincia", "").lower().replace(" ", "-")
    delegacion = datos.get("delegacion", "")
    cod_mun    = datos.get("cod_municipio", "")
    refcat     = datos.get("referencia_catastral", "")

    if lat and lon:
        link_gmaps = f"https://www.google.com/maps?q={lat},{lon}"
    else:
        direccion = " ".join(filter(None, [
            datos.get("tipo_via", ""), datos.get("nombre_via", ""),
            datos.get("numero", ""), datos.get("municipio", ""),
        ]))
        link_gmaps = f"https://www.google.com/maps/search/{quote(direccion)}"

    if cp:
        link_idealista = f"https://www.idealista.com/geo/venta-viviendas/codigo-postal-{cp}/"
    elif lat and lon:
        link_idealista = f"https://www.idealista.com/maps/venta-viviendas/?center={lat},{lon}&zoom=19"
    else:
        link_idealista = f"https://www.idealista.com/geo/venta-viviendas/{mun}-{prov}/"

    if delegacion and cod_mun:
        link_catastro = (
            f"https://www1.sedecatastro.gob.es/CYCBienInmueble/OVCConCiud.aspx"
            f"?del={delegacion}&mun={cod_mun}&RefC={refcat}&esBice=&RCBice1=&RCBice2="
            f"&DenoBice=&from=OVCBusqueda&pest=rc&RCCompleta={refcat}&final=&urlBuzon="
        )
    else:
        link_catastro = ""

    return link_gmaps, link_idealista, link_catastro


def obtener_plantas_edificio(refcat):
    """Consulta el numero de plantas del edificio usando la referencia del inmueble colectivo (14 chars).
    Hace una segunda llamada al catastro con la referencia de edificio en lugar de la unidad."""
    ref_edificio = refcat[:14]
    try:
        params = urlencode({"Provincia": "", "Municipio": "", "RC": ref_edificio})
        root = get_xml(f"{BASE}/OVCCallejero.asmx/Consulta_DNPRC?{params}")
        npt = texto(root, "npt")
        return npt
    except RuntimeError:
        return ""


def consultar_refcat(refcat_input):
    refcat = limpiar_refcat(refcat_input)
    resultado = {"referencia_catastral": refcat}
    try:
        resultado.update(obtener_datos_callejero(refcat))
    except RuntimeError as e:
        resultado["error_callejero"] = str(e)
    try:
        resultado.update(obtener_coordenadas(refcat))
    except RuntimeError:
        pass
    # Segunda llamada para obtener numero de plantas del edificio
    if not resultado.get("num_plantas"):
        plantas = obtener_plantas_edificio(refcat)
        if plantas:
            resultado["num_plantas"] = plantas
    return resultado


def detectar_planta_baja(datos):
    """Devuelve aviso si el inmueble esta en planta baja, sotano o semisotano."""
    planta = datos.get("planta", "").strip().upper()
    if not planta:
        return ""
    PLANTAS_BAJAS  = {"00", "PB", "B", "BJ", "BAJA", "0", "PT"}
    PLANTAS_SOTANO = {"SS", "SO", "ST", "SB", "S1", "S2", "S3", "-1", "-2", "SO1", "SO2"}
    if planta in PLANTAS_BAJAS:
        return "Planta baja (-10/15%)"
    if planta in PLANTAS_SOTANO:
        return "Sotano (-20/25%)"
    return ""


# ---------------------------------------------------------------------------
# Utilidades Idealista
# ---------------------------------------------------------------------------

def shape_circulo(lat, lon, radio_m=200, puntos=12):
    """Genera un poligono circular codificado en polyline (12 puntos = mejor aproximacion)."""
    import math
    import polyline as pl
    coords = []
    for i in range(puntos + 1):
        angulo = 2 * math.pi * i / puntos
        dlat = (radio_m / 111320) * math.cos(angulo)
        dlon = (radio_m / (111320 * math.cos(math.radians(float(lat))))) * math.sin(angulo)
        coords.append((round(float(lat) + dlat, 5), round(float(lon) + dlon, 5)))
    encoded = pl.encode(coords, 5)
    return quote(f"(({encoded}))")


def filtrar_outliers_iqr(pm2, pl, links):
    """Elimina outliers de precio/m2 con el metodo IQR. Mantiene indices sincronizados."""
    if len(pm2) < 4:
        return pm2, pl, links
    sorted_vals = sorted(pm2)
    n = len(sorted_vals)
    q1 = sorted_vals[n // 4]
    q3 = sorted_vals[(3 * n) // 4]
    iqr = q3 - q1
    if iqr == 0:
        return pm2, pl, links
    lower = q1 - 1.5 * iqr
    upper = q3 + 1.5 * iqr
    filtrado = [(p, l, lk) for p, l, lk in zip(pm2, pl, links) if lower <= p <= upper]
    if not filtrado:
        return pm2, pl, links  # fallback: no filtrar si elimina todo
    p_f, l_f, lk_f = zip(*filtrado)
    return list(p_f), list(l_f), list(lk_f)


def sacar_precio_idealista(driver, lat, lon, cp, mun, prov, tipo_manual="", superficie=""):
    from bs4 import BeautifulSoup
    import time
    import random

    t = tipo_manual.lower().strip()
    es_parking = t in ("parking", "garaje", "garajes")
    es_local   = t in ("local", "locales")

    # -----------------------------------------------------------------------
    def _parsear_anuncios(anuncios, parking, vistos):
        """Extrae precios/m2 de una lista de articulos de Idealista."""
        pm2, pl, links = [], [], []
        m2_min_filter = 5 if parking else 20
        for a in anuncios:
            precio  = a.find("span", class_="item-price")
            detalles = a.find_all("span", class_="item-detail")
            link_tag = a.find("a", class_="item-link")
            if not precio:
                continue
            precio_num = re.sub(r"[^0-9]", "", precio.text.strip())
            m2_num = None
            for d in detalles:
                m = re.search(r"([0-9]+)\s*m", d.text)
                if m:
                    m2_num = int(m.group(1))
                    break
            link = (f"https://www.idealista.com{link_tag['href']}"
                    if link_tag and link_tag.get("href") else "")
            if precio_num and parking and not m2_num:
                clave = f"{precio_num}_0"
                if clave not in vistos:
                    vistos.add(clave)
                    pm2.append(int(precio_num))
                    pl.append(int(precio_num))
                    links.append(link)
            elif precio_num and m2_num and m2_num > m2_min_filter:
                clave = f"{precio_num}_{m2_num}"
                if clave in vistos:
                    continue
                vistos.add(clave)
                pm2.append(int(precio_num) / m2_num)
                pl.append(int(precio_num))
                links.append(link)
        return pm2, pl, links

    # -----------------------------------------------------------------------
    def buscar_url(url_base, radio_label, parking=False):
        """Carga hasta 3 paginas de Idealista hasta obtener >=15 comparables."""
        pm2_total, pl_total, links_total = [], [], []
        vistos = set()

        for pagina in range(1, 4):
            if pagina == 1:
                url_pag = url_base
            else:
                sep = "&" if "?" in url_base else "?"
                url_pag = f"{url_base}{sep}pagina={pagina}"

            print(f"  URL ({radio_label}, pag {pagina}): {url_pag}")
            driver.get(url_pag)
            time.sleep(random.uniform(8, 12))
            driver.execute_script("window.scrollTo(0, 500)")
            time.sleep(random.uniform(3, 5))
            soup = BeautifulSoup(driver.page_source, "html.parser")

            anuncios = [a for a in soup.find_all("article") if "item" in a.get("class", [])]
            if not anuncios:
                break

            pm2_pag, pl_pag, links_pag = _parsear_anuncios(anuncios, parking, vistos)
            pm2_total.extend(pm2_pag)
            pl_total.extend(pl_pag)
            links_total.extend(links_pag)

            if len(pm2_total) >= 15:
                break

            # Comprobar si hay pagina siguiente
            sig = soup.find("a", class_=re.compile(r"icon-arrow-right-after"))
            if not sig:
                break

        return pm2_total, pl_total, links_total

    # -----------------------------------------------------------------------
    def calcular_resultado(pm2, pl, links, tipo, radio):
        """Aplica IQR, calcula mediana y coeficiente de variacion."""
        pm2, pl, links = filtrar_outliers_iqr(pm2, pl, links)
        if not pm2:
            return None
        precio_medio = round(statistics.median(pm2))
        precio_min_v = min(pm2)
        precio_max_v = max(pm2)
        idx_min = pm2.index(precio_min_v)
        idx_max = pm2.index(precio_max_v)
        coef_var = (round(statistics.stdev(pm2) / statistics.mean(pm2) * 100, 1)
                    if len(pm2) > 1 else 0.0)
        return {
            "precio_m2_medio":  precio_medio,
            "precio_m2_min":    round(precio_min_v),
            "precio_m2_max":    round(precio_max_v),
            "precio_min":       pl[idx_min],
            "precio_max":       pl[idx_max],
            "link_min":         links[idx_min],
            "link_max":         links[idx_max],
            "num_comparables":  len(pm2),
            "coef_variacion":   coef_var,
            "tipo":             tipo,
            "radio":            radio,
        }

    # -----------------------------------------------------------------------
    # PARKING
    # -----------------------------------------------------------------------
    if es_parking:
        tipo = "garajes"
        if lat and lon:
            shape = shape_circulo(lat, lon, radio_m=200)
            url = f"https://www.idealista.com/areas/venta-garajes/?shape={shape}"
        elif cp:
            url = f"https://www.idealista.com/venta-garajes/{mun}-{prov}/con-codigos_postales-{cp}/"
        else:
            url = f"https://www.idealista.com/venta-garajes/{mun}-{prov}/"

        radio_usado = "200m"
        pm2, pl, links = buscar_url(url, "200m", parking=True)

        if len(pm2) < 5 and lat and lon:
            print(f"  Solo {len(pm2)} comparables parking, ampliando a 500m...")
            shape2 = shape_circulo(lat, lon, radio_m=500)
            url2 = f"https://www.idealista.com/areas/venta-garajes/?shape={shape2}"
            pm2, pl, links = buscar_url(url2, "500m", parking=True)
            radio_usado = "500m"

        return calcular_resultado(pm2, pl, links, tipo, radio_usado)

    # -----------------------------------------------------------------------
    # LOCAL
    # -----------------------------------------------------------------------
    elif es_local:
        tipo = "locales"
        if lat and lon:
            shape = shape_circulo(lat, lon, radio_m=200)
            url = f"https://www.idealista.com/areas/venta-locales/?shape={shape}"
        elif cp:
            url = f"https://www.idealista.com/venta-locales/{mun}-{prov}/con-codigos_postales-{cp}/"
        else:
            url = f"https://www.idealista.com/venta-locales/{mun}-{prov}/"

        radio_usado = "200m"
        pm2, pl, links = buscar_url(url, "200m")

        if len(pm2) < 5 and lat and lon:
            print(f"  Solo {len(pm2)} comparables local, ampliando a 500m...")
            shape2 = shape_circulo(lat, lon, radio_m=500)
            url2 = f"https://www.idealista.com/areas/venta-locales/?shape={shape2}"
            pm2, pl, links = buscar_url(url2, "500m")
            radio_usado = "500m"

        return calcular_resultado(pm2, pl, links, tipo, radio_usado)

    # -----------------------------------------------------------------------
    # PISO / CHALET
    # -----------------------------------------------------------------------
    else:
        tipo = "chalets" if t in ("chalet", "chalets") else "pisos"

        filtro_m2 = ""
        if superficie:
            try:
                m2 = float(superficie.replace(",", "."))
                m2_min = max(20, round(m2 * 0.8))
                m2_max = round(m2 * 1.2)
                filtro_m2 = f"con-metros-cuadrados-mas-de_{m2_min},metros-cuadrados-menos-de_{m2_max},"
            except Exception:
                pass

        if lat and lon:
            shape = shape_circulo(lat, lon, radio_m=200)
            url = f"https://www.idealista.com/areas/venta-viviendas/{filtro_m2}sin-inquilinos/?shape={shape}"
        elif cp:
            url = f"https://www.idealista.com/geo/venta-viviendas/codigo-postal-{cp}/{filtro_m2}sin-inquilinos/"
        else:
            url = f"https://www.idealista.com/venta-viviendas/{mun}-{prov}/{filtro_m2}sin-inquilinos/"

        radio_usado = "200m"
        pm2, pl, links = buscar_url(url, "200m")

        if len(pm2) < 5 and lat and lon:
            print(f"  Solo {len(pm2)} comparables, ampliando a 500m...")
            shape2 = shape_circulo(lat, lon, radio_m=500)
            url2 = f"https://www.idealista.com/areas/venta-viviendas/{filtro_m2}sin-inquilinos/?shape={shape2}"
            pm2, pl, links = buscar_url(url2, "500m")
            radio_usado = "500m"

        return calcular_resultado(pm2, pl, links, tipo, radio_usado)


# ---------------------------------------------------------------------------
# Exportar Excel
# ---------------------------------------------------------------------------

def exportar_excel(lista_datos, filename="resultados_catastro.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Catastro AVM"

    header_font  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body_font    = Font(name="Arial", size=10)
    link_font    = Font(name="Arial", size=10, color="1155CC", underline="single")
    green_font   = Font(name="Arial", size=10, bold=True, color="1B5E20")
    header_fill  = PatternFill("solid", start_color="1F4E79")
    alt_fill     = PatternFill("solid", start_color="D6E4F0")
    green_fill   = PatternFill("solid", start_color="E8F5E9")
    green_hdr    = PatternFill("solid", start_color="2E7D32")
    warn_hdr     = PatternFill("solid", start_color="E65100")
    center       = Alignment(horizontal="center", vertical="center")
    left         = Alignment(horizontal="left",   vertical="center")
    thin         = Side(style="thin", color="BFBFBF")
    border       = Border(left=thin, right=thin, top=thin, bottom=thin)

    todas_columnas = COLUMNAS + COLUMNAS_LINKS
    total_cols = len(todas_columnas)

    # Fila 1: titulo
    ws.merge_cells(f"A1:{get_column_letter(total_cols)}1")
    titulo = ws["A1"]
    titulo.value     = f"Consulta Catastral AVM v5 — {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    titulo.font      = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    titulo.fill      = PatternFill("solid", start_color="0D3050")
    titulo.alignment = center
    ws.row_dimensions[1].height = 24

    # Fila 2: cabeceras
    for col_idx, (key, label) in enumerate(COLUMNAS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=label)
        if key in GREEN_KEYS:
            cell.fill = green_hdr
        elif key in YELLOW_KEYS:
            cell.fill = PatternFill("solid", start_color="F9A825")
        elif key in WARNING_KEYS:
            cell.fill = warn_hdr
        else:
            cell.fill = header_fill
        cell.font      = header_font
        cell.alignment = center
        cell.border    = border

    for col_idx, (_, label) in enumerate(COLUMNAS_LINKS, start=len(COLUMNAS) + 1):
        cell = ws.cell(row=2, column=col_idx, value=label)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = border

    ws.row_dimensions[2].height = 20

    # Filas de datos
    for fila_idx, datos in enumerate(lista_datos, start=3):
        is_alt = fila_idx % 2 == 0

        for col_idx, (key, _) in enumerate(COLUMNAS, start=1):
            valor = datos.get(key, "")
            cell  = ws.cell(row=fila_idx, column=col_idx, value=valor)
            cell.border    = border
            cell.alignment = center if col_idx == 1 else left

            if key == "coef_variacion":
                try:
                    cv = float(str(valor).replace(",", ".")) if valor else 0
                except Exception:
                    cv = 0
                if cv <= 15:
                    cell.fill = green_fill
                    cell.font = green_font
                elif cv <= 25:
                    cell.fill = PatternFill("solid", start_color="FFF9C4")
                    cell.font = Font(name="Arial", size=10, bold=True, color="7B5A00")
                else:
                    cell.fill = PatternFill("solid", start_color="FFE0B2")
                    cell.font = Font(name="Arial", size=10, bold=True, color="BF360C")

            elif key == "planta_baja_flag" and valor:
                cell.fill = PatternFill("solid", start_color="FFE0B2")
                cell.font = Font(name="Arial", size=10, bold=True, color="E65100")

            elif key in GREEN_KEYS:
                cell.font = green_font
                cell.fill = green_fill
                if key == "precio_min_comparable" and datos.get("link_min_comparable"):
                    cell.hyperlink = datos["link_min_comparable"]
                    cell.font = Font(name="Arial", size=10, bold=True, color="1155CC", underline="single")
                elif key == "precio_max_comparable" and datos.get("link_max_comparable"):
                    cell.hyperlink = datos["link_max_comparable"]
                    cell.font = Font(name="Arial", size=10, bold=True, color="1155CC", underline="single")

            elif key in YELLOW_KEYS:
                cell.font = Font(name="Arial", size=10, bold=True, color="7B5A00")
                cell.fill = (green_fill if valor == "200m"
                             else PatternFill("solid", start_color="FFF9C4"))

            else:
                cell.font = body_font
                cell.fill = alt_fill if is_alt else PatternFill(fill_type=None)

        link_gmaps, link_idealista, link_catastro = generar_links(datos)
        offset = len(COLUMNAS) + 1
        for i, (txt, url) in enumerate([
            ("Ver en Maps",      link_gmaps),
            ("Ver en Idealista", link_idealista),
            ("Ver en Catastro",  link_catastro),
        ]):
            cell = ws.cell(row=fila_idx, column=offset + i, value=txt)
            if url:
                cell.hyperlink = url
            cell.font      = link_font
            cell.alignment = center
            cell.border    = border
            if is_alt:
                cell.fill = alt_fill

        ws.row_dimensions[fila_idx].height = 18

    # Anchos de columna (49 COLUMNAS + 3 LINKS = 52)
    anchos = [
        22, 8, 28, 8, 8, 8, 8, 8, 9, 20, 16, 22, 14, 16, 14, 14, 14, 12, 12, 12,
        16, 12, 12, 12, 14, 12, 12, 12, 16, 10, 20, 22, 22, 22, 12, 12, 14, 18, 16, 16,
        16, 22, 20, 22, 22, 14, 13, 20, 25, 16, 18, 16,
    ]
    for i, ancho in enumerate(anchos[:total_cols], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(total_cols)}{len(lista_datos) + 2}"
    wb.save(filename)
    return filename


# ---------------------------------------------------------------------------
# Consola
# ---------------------------------------------------------------------------

def imprimir_resultado(datos):
    print(f"\n{'─' * 55}")
    for key, label in COLUMNAS:
        valor = datos.get(key)
        if valor:
            print(f"  {label:<38} {valor}")
    print(f"{'─' * 55}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    from openpyxl import load_workbook as _lw

    if len(sys.argv) < 2:
        print("Uso: python catastro_avm_v5.py referencias.xlsx")
        sys.exit(1)

    referencias = []
    for arg in sys.argv[1:]:
        if ".xls" in arg.lower():
            try:
                wb = _lw(arg)
                ws = wb.active
                col_ref  = 1
                col_tipo = None
                for cell in ws[1]:
                    if cell.value and "catastral" in str(cell.value).lower():
                        col_ref = cell.column
                    if cell.value and "tipo" in str(cell.value).lower():
                        col_tipo = cell.column
                count = 0
                for row in ws.iter_rows(min_row=2):
                    ref_val = str(row[col_ref - 1].value).strip() if row[col_ref - 1].value else ""
                    if not ref_val or ref_val.lower() == "none":
                        continue
                    tipo_val = ""
                    if col_tipo:
                        tipo_val = str(row[col_tipo - 1].value).strip() if row[col_tipo - 1].value else ""
                        if tipo_val.lower() == "none":
                            tipo_val = ""
                    referencias.append((ref_val, tipo_val))
                    count += 1
                print(f"Leidas {count} referencias de {arg}")
            except Exception as e:
                print(f"ERROR: {e}")
        else:
            referencias.append((arg, ""))

    if not referencias:
        print("No se encontraron referencias.")
        sys.exit(1)

    print(f"\nTotal: {len(referencias)} referencias")
    resp = input("\nSacar precio de mercado de Idealista? (s/n): ").strip().lower()
    usar_idealista = resp == "s"

    driver = None
    if usar_idealista:
        try:
            import undetected_chromedriver as uc
            import time
            print("Abriendo Chrome...")
            driver = uc.Chrome(version_main=146)
            time.sleep(3)
            driver.get("https://www.idealista.com")
            time.sleep(5)
            print("Chrome listo.")
        except Exception as e:
            print(f"Error abriendo Chrome: {e}")
            driver = None

    resultados = []
    for ref, tipo_manual in referencias:
        print(f"\nConsultando: {limpiar_refcat(ref)} ...")
        datos = consultar_refcat(ref)

        # Autodeteccion de tipo si no viene del Excel
        if not tipo_manual:
            uso = datos.get("uso_principal", "").upper()
            for uso_key, tipo_det in USO_A_TIPO.items():
                if uso_key in uso:
                    tipo_manual = tipo_det
                    print(f"  Tipo autodetectado: {tipo_manual} (uso_principal={uso})")
                    break

        # Flag planta baja / sotano
        flag = detectar_planta_baja(datos)
        if flag:
            datos["planta_baja_flag"] = flag
            print(f"  Aviso: {flag}")

        if usar_idealista and driver:
            sup  = datos.get("m2_vivienda") or datos.get("superficie_m2", "")
            cp   = datos.get("codigo_postal", "")
            mun  = datos.get("municipio", "").lower().replace(" ", "-")
            prov = datos.get("provincia", "").lower().replace(" ", "-")
            try:
                res = sacar_precio_idealista(
                    driver,
                    datos.get("coord_lat_wgs84", ""),
                    datos.get("coord_lon_wgs84", ""),
                    cp, mun, prov,
                    tipo_manual=tipo_manual,
                    superficie=sup,
                )
                if res:
                    es_parking_res = res["tipo"] == "garajes"
                    datos["tipo_busqueda"]         = res["tipo"]
                    datos["precio_m2_idealista"]   = str(res["precio_m2_medio"])
                    datos["precio_m2_min"]         = str(res["precio_m2_min"])
                    datos["precio_m2_max"]         = str(res["precio_m2_max"])
                    datos["coef_variacion"]        = str(res["coef_variacion"])
                    datos["num_comparables"]       = str(res["num_comparables"])
                    datos["radio_busqueda"]        = res.get("radio", "200m")
                    datos["precio_min_comparable"] = f"{res['precio_min']:,} EUR".replace(",", ".")
                    datos["precio_max_comparable"] = f"{res['precio_max']:,} EUR".replace(",", ".")
                    datos["link_min_comparable"]   = res.get("link_min", "")
                    datos["link_max_comparable"]   = res.get("link_max", "")

                    if es_parking_res:
                        datos["valoracion_idealista"] = f"{res['precio_m2_medio']:,} EUR".replace(",", ".")
                    elif sup:
                        try:
                            val = round(res["precio_m2_medio"] * float(sup.replace(",", ".")))
                            datos["valoracion_idealista"] = f"{val:,} EUR".replace(",", ".")
                        except Exception:
                            pass

                    m2_viv = datos.get("m2_vivienda", "")
                    if m2_viv and not es_parking_res:
                        try:
                            val_viv = round(res["precio_m2_medio"] * float(m2_viv.replace(",", ".")))
                            datos["valoracion_vivienda"] = f"{val_viv:,} EUR".replace(",", ".")
                        except Exception:
                            pass

                    fiab = ("FIABLE" if res["coef_variacion"] <= 15
                            else "MODERADO" if res["coef_variacion"] <= 25
                            else "REVISAR")
                    print(f"  Precio m2: {res['precio_m2_medio']} EUR | "
                          f"CV: {res['coef_variacion']}% ({fiab}) | "
                          f"{res['num_comparables']} comp | {res['tipo']} | {res['radio']}")
            except Exception as e:
                print(f"  Error Idealista: {e}")

        imprimir_resultado(datos)
        resultados.append(datos)

    if driver:
        try:
            driver.quit()
        except Exception:
            pass

    fichero = exportar_excel(resultados)
    print(f"\nExcel guardado en: {fichero}\n")
    try:
        os.startfile(fichero)
    except Exception:
        pass
